from __future__ import annotations

import io
import re
from collections import Counter, defaultdict
from datetime import datetime

import pandas as pd
import pdfplumber
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "output"
CURRENT_YEAR_PATTERN = None  # inferred from the USN year tokens in the source PDF
CURRENT_ANALYSIS = None

SUBJECTS = [
    {"code": "BCOM5-DSC1", "name": "Financial Management"},
    {"code": "BCOM5-DSC2", "name": "Income Tax Law and Practice-I"},
    {"code": "BCOM5-DSC3", "name": "Principles and Practice of Auditing"},
    {"code": "BCOM5-DSE-A1", "name": "Indian Accounting Standards-I"},
    {"code": "BCOM5-DSE-F1", "name": "Financial Institutions and Markets"},
    {"code": "BCOM5-VOC1", "name": "GST-Law and Practice"},
    {"code": "COM5ES-SEC-2", "name": "Employability Skill"},
]
STATUS_WORDS = {"PASS", "FAIL", "ABSENT", "P", "F", "A"}
USN_RE = re.compile(r"\bU\d{2}[A-Z]{1,5}\d{2}[A-Z0-9]{3,}\b", re.I)
NUMBER_RE = re.compile(r"\b\d{1,3}\b")

app = Flask(__name__)


def clean_name(value: str) -> str:
    value = re.sub(r"\s+", " ", value).strip(" -|:")
    return value.title() if value.isupper() else value


def status(value: str) -> str:
    value = value.upper()
    return {"P": "PASS", "F": "FAIL", "A": "ABSENT"}.get(value, value)


def parse_subject_row(text: str, code: str) -> dict | None:
    match = re.search(re.escape(code), text, re.I)
    if not match:
        return None
    next_code = re.search(r"(?:BCOM5-DSC1|BCOM5-DSC2|BCOM5-DSC3|BCOM5-DSE-A1|BCOM5-DSE-F1|BCOM5-VOC1|COM5ES-SEC-2)", text[match.end():], re.I)
    tail = text[match.end():match.end() + next_code.start()] if next_code else text[match.end():]
    words = re.findall(r"[A-Za-z]+", tail.upper())
    found_status = next((status(word) for word in words if word in STATUS_WORDS), None)
    nums = [int(item) for item in NUMBER_RE.findall(tail)]
    if not nums and not found_status:
        return None
    theory = internal = total = None
    if len(nums) >= 3:
        theory, internal, total = nums[-3:]
    elif len(nums) == 2:
        theory, total = nums[-2:]
    elif len(nums) == 1:
        total = nums[0]
    if found_status is None:
        found_status = "PASS" if total is not None else "ABSENT"
    return {"theory": theory, "internal": internal, "total": total, "status": found_status}


def parse_pdf(pdf_bytes: bytes) -> tuple[list[dict], list[str]]:
    if not pdf_bytes:
        raise ValueError("The uploaded PDF is empty.")
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        pages = [page.extract_text(x_tolerance=2, y_tolerance=3) or "" for page in pdf.pages]
    text = "\n".join(pages)
    if not text.strip():
        raise ValueError("The PDF contains no readable text.")

    record_starts = list(re.finditer(r"(?m)^\d{5}\s+Sub\b", text))
    records, warnings = [], []
    for index, start in enumerate(record_starts):
        block = text[start.start():record_starts[index + 1].start() if index + 1 < len(record_starts) else len(text)]
        usn_match = USN_RE.search(block)
        if not usn_match:
            continue
        usn = usn_match.group(0).upper()
        before_usn = re.sub(r"Total", "", block[:usn_match.start()], flags=re.I)
        before_usn = re.sub(r"\s+", " ", before_usn)
        code_matches = re.findall(r"(?:DSC1|DSC2|DSC3|DSE-A1|DSE-F1|VOC1|SEC\s*-\s*2)", before_usn, re.I)
        code_order = []
        for code in code_matches:
            suffix = re.sub(r"\s+", "", code).upper()
            normalized = "COM5ES-SEC-2" if suffix == "SEC-2" else f"BCOM5-{suffix}"
            if normalized not in code_order:
                code_order.append(normalized)
        marks_line = next((line for line in block[usn_match.end():].splitlines() if "Th" in line), "")
        mark_tokens = re.findall(r"\d{1,3}\s*\+\s*\d{1,3}|-\s*-", marks_line)
        mark_pairs = []
        for token in mark_tokens:
            values = re.findall(r"\d{1,3}", token)
            mark_pairs.append((int(values[0]), int(values[1])) if len(values) == 2 else (None, None))
        total_match = re.search(r"(?m)^Result:\s*(PASS|FAIL|ABSENT)\s+Total\s+(.+)$", block, re.I)
        subject_totals = []
        overall_result = None
        if total_match:
            overall_result = total_match.group(1).upper()
            total_values = re.findall(r"\d+(?:\.\d+)?", total_match.group(2))
            subject_totals = [float(value) for value in total_values]
        status_tail = block[block.rfind("M.C.No"):] if "M.C.No" in block else ""
        subject_statuses = [status(word) for word in re.findall(r"\b(PASS|FAIL|ABSENT|P|F|A)\b", status_tail, re.I)]
        subjects = {}
        for position, code in enumerate(code_order):
            pair = mark_pairs[position] if position < len(mark_pairs) else (None, None)
            raw_total = subject_totals[position] if position < max(len(code_order), len(subject_totals)) else None
            subject_total = int(raw_total) if raw_total is not None and raw_total.is_integer() else raw_total
            subject_status = subject_statuses[position] if position < len(subject_statuses) else ("ABSENT" if pair == (None, None) else "PASS")
            subjects[code] = {"theory": pair[0], "internal": pair[1], "total": subject_total, "status": subject_status}
            if pair[0] is not None and pair[1] is not None and subject_total is not None and pair[0] + pair[1] != subject_total:
                warnings.append(f"{usn}: {code} theory + internal mismatch")
        grand_total = None
        if subject_totals:
            candidate = subject_totals[-1] if len(subject_totals) > len(code_order) else (sum(subject_totals) if len(subject_totals) == len(code_order) else None)
            grand_total = int(candidate) if candidate is not None and candidate.is_integer() else candidate
        if grand_total is not None and len(subject_totals) > len(code_order):
            expected = sum(subject_totals[:-1])
            if round(expected, 2) != round(grand_total, 2):
                warnings.append(f"{usn}: subject totals do not equal grand total")
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        usn_index = next((line_index for line_index, line in enumerate(lines) if usn in line), 0)
        name_line = lines[usn_index + 1] if usn_index + 1 < len(lines) else "Unknown"
        name_line = re.sub(r"^(?:NAME|STUDENT NAME)\s*[:.-]?\s*", "", name_line, flags=re.I)
        sgpa_match = re.search(r"SGPA\s+([\d.]+)", block, re.I)
        cgpa_match = re.search(r"CGPA\s+([\d.]+)", block, re.I)
        grade_match = re.search(r"(?:Class\s+([^\n]+?)\s+Max\.|Term Grade:\s*([^\n]+))", block, re.I)
        percentage = round(grand_total / 700 * 100, 2) if grand_total is not None and len(subjects) == 7 and all(item.get("total") is not None for item in subjects.values()) else None
        result = overall_result or ("ABSENT" if subjects and all(item["status"] == "ABSENT" for item in subjects.values()) else ("PASS" if subjects and all(item["status"] == "PASS" for item in subjects.values()) else "FAIL"))
        records.append({"usn": usn, "name": clean_name(name_line), "subjects": subjects, "total": grand_total, "percentage": percentage, "result": result, "sgpa": sgpa_match.group(1) if sgpa_match else None, "cgpa": cgpa_match.group(1) if cgpa_match else None, "grade": (grade_match.group(1) or grade_match.group(2)).strip() if grade_match else None})
    infer_year_pattern(records)
    for record in records:
        record["year_type"] = "Current Year" if CURRENT_YEAR_PATTERN and CURRENT_YEAR_PATTERN.search(record["usn"]) else "Backlog / Repeater"
    if not records:
        raise ValueError("No student records were found in the uploaded PDF.")
    return records, warnings


def infer_year_pattern(records: list[dict]) -> None:
    global CURRENT_YEAR_PATTERN
    years = Counter(
        match.group(1)
        for record in records
        for match in [re.search(r"(?:U\d{2}[A-Z]{1,5})(\d{2})", record["usn"], re.I)]
        if match
    )
    if not years:
        CURRENT_YEAR_PATTERN = None
        return
    current_year = years.most_common(1)[0][0]
    CURRENT_YEAR_PATTERN = re.compile(rf"^U\d{{2}}[A-Z]{{1,5}}{current_year}", re.I)


def subject_stats(records: list[dict]) -> list[dict]:
    result = []
    for subject in SUBJECTS:
        values = [r["subjects"].get(subject["code"]) for r in records]
        numeric = [v["total"] for v in values if v and v.get("total") is not None]
        passed = sum(bool(v and v["status"] == "PASS") for v in values)
        failed = sum(bool(v and v["status"] == "FAIL") for v in values)
        absent = sum(bool(v and v["status"] == "ABSENT") for v in values)
        appeared = passed + failed
        result.append({**subject, "appeared": appeared, "passed": passed, "failed": failed, "absent": absent, "pass_percentage": round(passed / appeared * 100, 2) if appeared else 0, "fail_percentage": round(failed / appeared * 100, 2) if appeared else 0, "average": round(sum(numeric) / len(numeric), 2) if numeric else 0, "highest": max(numeric) if numeric else 0, "lowest": min(numeric) if numeric else 0})
    return result


def build_view_data(records: list[dict], warnings: list[str], filename: str = "") -> dict:
    attended = sum(r["result"] in {"PASS", "FAIL"} for r in records)
    passed = sum(r["result"] == "PASS" for r in records)
    failed = sum(r["result"] == "FAIL" for r in records)
    absent = sum(r["result"] == "ABSENT" for r in records)
    current = [r for r in records if r["year_type"] == "Current Year"]
    backlog = [r for r in records if r["year_type"] != "Current Year"]
    stats = subject_stats(records)
    subject_failures = {s["code"]: [{"usn": r["usn"], "name": r["name"], **r["subjects"].get(s["code"], {})} for r in records if r["subjects"].get(s["code"], {}).get("status") == "FAIL"] for s in SUBJECTS}
    multiple = []
    for r in records:
        failed_subjects = [s["name"] for s in SUBJECTS if r["subjects"].get(s["code"], {}).get("status") == "FAIL"]
        if len(failed_subjects) > 1:
            multiple.append({"usn": r["usn"], "name": r["name"], "count": len(failed_subjects), "subjects": failed_subjects})
    eligible = sorted([r for r in current if r["result"] == "PASS" and r["percentage"] is not None and all(r["subjects"].get(s["code"], {}).get("status") == "PASS" for s in SUBJECTS)], key=lambda r: (r["total"] or 0, r["percentage"] or 0), reverse=True)
    toppers = [{"usn": r["usn"], "name": r["name"], "total": r["total"], "percentage": r["percentage"]} for r in eligible[:10]]
    subject_toppers = []
    for subject in SUBJECTS:
        top = sorted([r for r in current if r["subjects"].get(subject["code"], {}).get("status") == "PASS"], key=lambda r: r["subjects"][subject["code"]].get("total") or 0, reverse=True)[:3]
        subject_toppers.append({**subject, "toppers": [{"usn": r["usn"], "name": r["name"], **r["subjects"][subject["code"]]} for r in top]})
    serial = []
    for r in sorted(records, key=lambda x: (x["total"] or 0), reverse=True):
        serial.append({"usn": r["usn"], "name": r["name"], "total": r["total"], "percentage": r["percentage"], "result": r["result"], "year_type": r["year_type"], "sgpa": r["sgpa"], "cgpa": r["cgpa"], "grade": r["grade"], "subjects": r["subjects"]})
    current_percentages = [r["percentage"] for r in current if r["percentage"] is not None]
    summary = {"total": len(records), "attended": attended, "passed": passed, "failed": failed, "absent": absent, "pass_percentage": round(passed / attended * 100, 2) if attended else 0, "fail_percentage": round(failed / attended * 100, 2) if attended else 0, "current_year": len(current), "backlog": len(backlog)}
    current_summary = {"total": len(current), "attended": sum(r["result"] in {"PASS", "FAIL"} for r in current), "passed": sum(r["result"] == "PASS" for r in current), "failed": sum(r["result"] != "PASS" for r in current), "pass_percentage": round(sum(r["result"] == "PASS" for r in current) / sum(r["result"] in {"PASS", "FAIL"} for r in current) * 100, 2) if current else 0, "average_percentage": round(sum(current_percentages) / len(current_percentages), 2) if current_percentages else 0}
    return {"summary": summary, "current": current_summary, "subjects": stats, "students": serial, "subject_failures": subject_failures, "multiple_failures": multiple, "toppers": toppers, "subject_toppers": subject_toppers, "backlog_students": [{"usn": r["usn"], "name": r["name"], "result": r["result"], "percentage": r["percentage"] or "-"} for r in backlog], "warnings": warnings, "filename": filename}


def flatten_records(records: list[dict]) -> pd.DataFrame:
    rows = []
    for rank, record in enumerate(sorted(records, key=lambda x: x["total"] or 0, reverse=True), 1):
        row = {"Rank": rank, "Register Number": record["usn"], "Student Name": record["name"]}
        for subject in SUBJECTS:
            item = record["subjects"].get(subject["code"], {})
            prefix = subject["name"]
            row.update({f"{prefix} Theory": item.get("theory"), f"{prefix} Internal": item.get("internal"), f"{prefix} Total": item.get("total"), f"{prefix} Status": item.get("status", "ABSENT")})
        row.update({"Grand Total": record["total"], "Maximum Marks": 700, "Percentage": record["percentage"], "Result": record["result"], "SGPA": record["sgpa"], "CGPA": record["cgpa"], "Grade/Class": record["grade"], "Year Type": record["year_type"]})
        rows.append(row)
    return pd.DataFrame(rows)


def create_excel(data: dict, records: list[dict]) -> io.BytesIO:
    output = io.BytesIO()
    attended = data["summary"]["attended"]
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([{"Exam": "December 2025", "Semester": "V", "Course": "B.Com", "Total Students": data["summary"]["total"], "Attended": attended, "Absent": data["summary"]["total"] - attended, "Passed": data["summary"]["passed"], "Failed": data["summary"]["failed"], "Pass Percentage": data["summary"]["pass_percentage"], "Fail Percentage": data["summary"]["fail_percentage"], "Current Year Students": data["summary"]["current_year"], "Backlog Students": data["summary"]["backlog"]}]).to_excel(writer, sheet_name="Summary", index=False)
        flatten_records(records).to_excel(writer, sheet_name="Student Results", index=False)
        pd.DataFrame(data["subjects"]).rename(columns={"name": "Subject", "pass_percentage": "Pass %", "fail_percentage": "Fail %"}).drop(columns=["code"], errors="ignore").to_excel(writer, sheet_name="Subject Analysis", index=False)
        failed_rows = []
        for subject in SUBJECTS:
            for item in data["subject_failures"][subject["code"]]:
                failed_rows.append({"Subject": subject["name"], **item})
        pd.DataFrame(failed_rows).to_excel(writer, sheet_name="Failed Students", index=False)
        pd.DataFrame(data["toppers"]).to_excel(writer, sheet_name="Overall Toppers", index=False)
        subject_rows = [{"Subject": s["name"], "Rank": i, **t} for s in data["subject_toppers"] for i, t in enumerate(s["toppers"], 1)]
        pd.DataFrame(subject_rows).to_excel(writer, sheet_name="Subject Toppers", index=False)
        pd.DataFrame([data["current"]]).to_excel(writer, sheet_name="Current Year Analysis", index=False)
        pd.DataFrame(data["backlog_students"]).to_excel(writer, sheet_name="Backlog Analysis", index=False)
        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="172554")
        thin = Side(style="thin", color="D8DEE9")
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            for row in sheet.iter_rows():
                for cell in row:
                    cell.border = Border(bottom=thin)
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in column) + 2, 12), 34)
            for cell in sheet[1]:
                if "Percentage" in str(cell.value) or "%" in str(cell.value):
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row=row, column=cell.column).number_format = "0.00"
            status_fills = {"PASS": PatternFill("solid", fgColor="C6EFCE"), "FAIL": PatternFill("solid", fgColor="FFC7CE"), "ABSENT": PatternFill("solid", fgColor="FFEB9C")}
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if str(cell.value).upper() in status_fills:
                        cell.fill = status_fills[str(cell.value).upper()]
    output.seek(0)
    return output


@app.route("/")
def index():
    if CURRENT_ANALYSIS is None:
        return render_template("index.html", analysis_ready=False, data={})
    return render_template("index.html", analysis_ready=True, **CURRENT_ANALYSIS, data=CURRENT_ANALYSIS)


@app.route("/upload", methods=["POST"])
def upload():
    global CURRENT_ANALYSIS
    uploaded = request.files.get("file")
    if uploaded is None or not uploaded.filename:
        return jsonify(error="Please select a result PDF first."), 400
    if not uploaded.filename.lower().endswith(".pdf"):
        return jsonify(error="Invalid PDF. Please upload a valid result PDF."), 400
    try:
        pdf_bytes = uploaded.read()
        with pdfplumber.open(io.BytesIO(pdf_bytes)):
            pass
        records, warnings = parse_pdf(pdf_bytes)
        CURRENT_ANALYSIS = build_view_data(records, warnings, uploaded.filename)
        return jsonify(ok=True, filename=uploaded.filename, summary=CURRENT_ANALYSIS["summary"])
    except Exception as error:
        app.logger.exception("Unable to analyze uploaded PDF")
        return jsonify(error="Unable to analyze this PDF. The file may not contain a supported result format."), 422


@app.route("/reset", methods=["POST"])
def reset():
    global CURRENT_ANALYSIS
    CURRENT_ANALYSIS = None
    return jsonify(ok=True)


@app.route("/download-excel")
def download_excel():
    if CURRENT_ANALYSIS is None:
        return jsonify(error="Please upload a result PDF first."), 400
    data = CURRENT_ANALYSIS
    records = [{"usn": s["usn"], "name": s["name"], "subjects": s["subjects"], "total": s["total"], "percentage": s["percentage"], "result": s["result"], "sgpa": s["sgpa"], "cgpa": s["cgpa"], "grade": s["grade"], "year_type": s["year_type"]} for s in data["students"]]
    report = create_excel(data, records)
    filename = OUTPUT_DIR / f"student_result_analysis_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    OUTPUT_DIR.mkdir(exist_ok=True)
    filename.write_bytes(report.getvalue())
    report.seek(0)
    return send_file(report, as_attachment=True, download_name=filename.name, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=5000)
