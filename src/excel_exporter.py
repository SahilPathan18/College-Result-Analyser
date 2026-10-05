"""
Excel Export Engine for College Result Analyser.
Generates university tabulation register analysis workbooks matching the official
Bangalore University departmental ledger format:
- Institutional & Examination Headers
- Multi-tier Subject Headers (Subject name, SEE/IA breakdown)
- Student Data Grid with automated Excel formulas for totals and percentages
- Visual pass/fail highlighting (soft green / soft red)
- Subject-wise Result Analysis table with dynamic pass rates
- Overall Class Result Summary with grade tier segmentation (Distinction, First Class, etc.)
- Secondary worksheet for Backlog / Repeater candidates
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Color palette
CLR_NAVY = "1F4E78"        # Main dark blue headers
CLR_BLUE = "2E75B6"        # Sub-headers (SEE, IA, Total)
CLR_PASS_BG = "C6EFCE"     # Soft green fill for passed subjects
CLR_PASS_FG = "006100"     # Dark green text
CLR_FAIL_BG = "FFC7CE"     # Soft red fill for failed subjects
CLR_FAIL_FG = "9C0006"     # Dark red text
CLR_BORDER = "B7B7B7"      # Clean thin grid border
CLR_WHITE = "FFFFFF"


def _get_short_code(code: str) -> str:
    m = re.search(r"-(.+)$", code)
    return m.group(1).strip() if m else code.strip()


def _format_semester(sem: str) -> str:
    sem_str = str(sem).strip().upper()
    num_map = {
        "1": "1st Sem", "2": "2nd Sem", "3": "3rd Sem", "4": "4th Sem", "5": "5th Sem", "6": "6th Sem", "7": "7th Sem", "8": "8th Sem",
        "I": "1st Sem", "II": "2nd Sem", "III": "3rd Sem", "IV": "4th Sem", "V": "5th Sem", "VI": "6th Sem", "VII": "7th Sem", "VIII": "8th Sem"
    }
    return num_map.get(sem_str, f"{sem_str} Sem" if "SEM" not in sem_str else sem_str)


def _format_program(prog: str) -> str:
    p = prog.strip()
    if "Computer Applications" in p:
        return "BCA"
    if "Business Administration" in p:
        return "BBA"
    if "Commerce" in p:
        return "B.Com"
    if "Science" in p:
        return "B.Sc"
    return p


def _is_practical(course: dict) -> bool:
    code = course.get("code", "").upper()
    name = course.get("name", "").upper()
    return "DSCP" in code or "LAB" in code or "LAB" in name or "PRACTICAL" in name


def _style_range(ws, cell_range, font=None, fill=None, border=None, alignment=None):
    """Ensure all cells within a merged range have consistent formatting and borders."""
    for row in ws[cell_range]:
        for cell in row:
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            if border:
                cell.border = border
            if alignment:
                cell.alignment = alignment


def _get_sem_ord(sem: str) -> str:
    sem_str = str(sem).strip().upper()
    num_map = {
        "1": "1st", "2": "2nd", "3": "3rd", "4": "4th", "5": "5th", "6": "6th", "7": "7th", "8": "8th",
        "I": "1st", "II": "2nd", "III": "3rd", "IV": "4th", "V": "5th", "VI": "6th", "VII": "7th", "VIII": "8th"
    }
    return num_map.get(sem_str, sem_str)


def export_pretty_excel(filepath: str, data: dict, df=None) -> None:
    """Generate the official institutional tabulation Excel workbook."""
    wb = openpyxl.Workbook()
    meta = data.get("metadata", {})
    courses = data.get("courses", [])
    all_students = data.get("students", [])

    current_students = sorted(
        [s for s in all_students if s.get("year_type") == "Current Year"],
        key=lambda s: s.get("usn", "")
    )
    if not current_students:
        current_students = sorted(all_students, key=lambda s: s.get("usn", ""))

    backlog_students = sorted(
        [s for s in all_students if s.get("year_type") != "Current Year"],
        key=lambda s: s.get("usn", "")
    )

    sem_ord = _get_sem_ord(meta.get("semester", "5th"))
    sheet_title = f"{sem_ord} Sem Result Analysis"
    ws = wb.active
    ws.title = sheet_title

    _render_analysis_sheet(ws, meta, courses, current_students, is_backlog=False)

    if backlog_students:
        ws_back = wb.create_sheet(title="Backlog Students")
        _render_analysis_sheet(ws_back, meta, courses, backlog_students, is_backlog=True)

    Path(filepath).parent.mkdir(parents=True, exist_ok=True)
    wb.save(filepath)


def _render_analysis_sheet(ws, meta: dict, courses: list[dict], students: list[dict], is_backlog=False) -> None:
    # Styles
    font_title = Font(name="Calibri", size=14, bold=True)
    font_subtitle = Font(name="Calibri", size=11, bold=True)
    font_sec_head = Font(name="Calibri", size=12, bold=True)
    font_th = Font(name="Calibri", size=10, bold=True, color=CLR_WHITE)
    font_sub_th = Font(name="Calibri", size=9, bold=True, color=CLR_WHITE)
    font_data = Font(name="Calibri", size=11, bold=False)
    font_data_bold = Font(name="Calibri", size=11, bold=True)
    font_pass = Font(name="Calibri", size=11, bold=True, color=CLR_PASS_FG)
    font_fail = Font(name="Calibri", size=11, bold=True, color=CLR_FAIL_FG)

    fill_navy = PatternFill("solid", fgColor=CLR_NAVY)
    fill_blue = PatternFill("solid", fgColor=CLR_BLUE)
    fill_pass = PatternFill("solid", fgColor=CLR_PASS_BG)
    fill_fail = PatternFill("solid", fgColor=CLR_FAIL_BG)

    thin_border = Border(
        left=Side(style="thin", color=CLR_BORDER),
        right=Side(style="thin", color=CLR_BORDER),
        top=Side(style="thin", color=CLR_BORDER),
        bottom=Side(style="thin", color=CLR_BORDER)
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    num_courses = len(courses)
    num_students = len(students)

    # Calculate column indices
    start_sub_col = 4
    end_sub_col = start_sub_col + (num_courses * 3) - 1
    col_tot_marks = end_sub_col + 1
    col_max_marks = col_tot_marks + 1
    col_pct = col_tot_marks + 2
    col_sgpa = col_tot_marks + 3
    col_cgpa = col_tot_marks + 4
    col_result = col_tot_marks + 5
    last_col = col_result
    last_col_letter = get_column_letter(last_col)

    # 1. Row 1: Institutional Header
    ws.row_dimensions[1].height = 24
    ws.merge_cells(f"A1:{last_col_letter}1")
    title_cell = ws["A1"]
    title_cell.value = meta.get("college", "ST PAULS COLLEGE, NELAGADERANAHALLI").upper()
    title_cell.font = font_title
    title_cell.alignment = align_center

    # 2. Row 2: Examination Subtitle
    ws.row_dimensions[2].height = 20
    ws.merge_cells(f"A2:{last_col_letter}2")
    sub_cell = ws["A2"]
    sem_ord = _get_sem_ord(meta.get("semester", "5th"))
    prog_name = _format_program(meta.get("program", "BCA"))
    month_name = meta.get("exam_month", "Dec 2025").replace("/", " ")
    uni_name = meta.get("university", "Bangalore University")
    cohort_tag = " (Backlog / Repeater Analysis)" if is_backlog else ""
    sub_cell.value = f"{sem_ord} Semester {prog_name} | Examination Result Analysis — {month_name} ({uni_name}){cohort_tag}"
    sub_cell.font = font_subtitle
    sub_cell.alignment = align_center

    # 3. Row Heights for Headers
    ws.row_dimensions[3].height = 28
    ws.row_dimensions[4].height = 22

    # 4. Multi-tier Column Headers
    # A3:A4: Sl No
    ws.merge_cells("A3:A4")
    ws["A3"].value = "Sl\nNo"
    _style_range(ws, "A3:A4", font=font_th, fill=fill_navy, border=thin_border, alignment=align_center)

    # B3:B4: Register No
    ws.merge_cells("B3:B4")
    ws["B3"].value = "Register No"
    _style_range(ws, "B3:B4", font=font_th, fill=fill_navy, border=thin_border, alignment=align_center)

    # C3:C4: Student Name
    ws.merge_cells("C3:C4")
    ws["C3"].value = "Student Name"
    _style_range(ws, "C3:C4", font=font_th, fill=fill_navy, border=thin_border, alignment=align_center)

    # Course Headers
    sub_total_col_letters = []
    for i, course in enumerate(courses):
        col1 = start_sub_col + (i * 3)
        col2 = col1 + 1
        col3 = col1 + 2
        let1, let2, let3 = get_column_letter(col1), get_column_letter(col2), get_column_letter(col3)
        sub_total_col_letters.append(let3)

        short_code = _get_short_code(course["code"])
        c_name = course.get("name", short_code)

        # Row 3: Subject Name (Code)
        ws.merge_cells(f"{let1}3:{let3}3")
        ws[f"{let1}3"].value = f"{c_name} ({short_code})"
        _style_range(ws, f"{let1}3:{let3}3", font=font_th, fill=fill_navy, border=thin_border, alignment=align_center)

        # Row 4: Column Sub-headers
        is_lab = _is_practical(course)
        sub_cols = [
            ("SEE(Pr)" if is_lab else "SEE", let1),
            ("IA(Pr)" if is_lab else "IA", let2),
            ("Total\n(SEE+IA)", let3)
        ]
        for label, ltr in sub_cols:
            c_sh = ws[f"{ltr}4"]
            c_sh.value = label
            c_sh.fill = fill_blue
            c_sh.font = font_sub_th
            c_sh.alignment = align_center
            c_sh.border = thin_border

    # Summary Headers (merged rows 3 to 4)
    sum_headers = [
        (col_tot_marks, "Total\nMarks\n(SEE+IA)"),
        (col_max_marks, "Max\nMarks"),
        (col_pct, "Overall %\n(Total/Max)"),
        (col_sgpa, "SGPA"),
        (col_cgpa, "CGPA"),
        (col_result, "Result"),
    ]
    for col_idx, h_title in sum_headers:
        ltr = get_column_letter(col_idx)
        ws.merge_cells(f"{ltr}3:{ltr}4")
        ws[f"{ltr}3"].value = h_title
        _style_range(ws, f"{ltr}3:{ltr}4", font=font_th, fill=fill_navy, border=thin_border, alignment=align_center)

    # 5. Student Rows
    start_row = 5
    for idx, s in enumerate(students):
        r = start_row + idx

        # Sl No
        ws.cell(r, 1, value=idx + 1).alignment = align_center
        ws.cell(r, 1).font = font_data
        ws.cell(r, 1).border = thin_border

        # Register No
        ws.cell(r, 2, value=s.get("usn", "")).alignment = align_center
        ws.cell(r, 2).font = font_data
        ws.cell(r, 2).border = thin_border

        # Student Name
        ws.cell(r, 3, value=s.get("name", "")).alignment = align_left
        ws.cell(r, 3).font = font_data
        ws.cell(r, 3).border = thin_border

        # Subject Marks
        for i, course in enumerate(courses):
            code = course["code"]
            if code not in s.get("subjects", {}):
                # Student did not take this elective/optional course
                col1 = start_sub_col + (i * 3)
                col2 = col1 + 1
                col3 = col1 + 2
                for ci in (col1, col2, col3):
                    c = ws.cell(r, ci, value="-")
                    c.alignment = align_center
                    c.font = font_data
                    c.border = thin_border
                continue

            sub_rec = s.get("subjects", {}).get(code, {})
            status = sub_rec.get("status", "PASS")
            failed_sub = status == "FAIL"

            fill_sub = fill_fail if failed_sub else fill_pass
            font_sub = font_fail if failed_sub else font_data

            col1 = start_sub_col + (i * 3)
            col2 = col1 + 1
            col3 = col1 + 2
            let1, let2 = get_column_letter(col1), get_column_letter(col2)

            see_val = sub_rec.get("theory")
            ia_val = sub_rec.get("internal")

            c1 = ws.cell(r, col1, value=see_val if see_val is not None else "")
            c2 = ws.cell(r, col2, value=ia_val if ia_val is not None else "")
            c3 = ws.cell(r, col3, value=f"={let1}{r}+{let2}{r}")

            for c in (c1, c2, c3):
                c.alignment = align_center
                c.fill = fill_sub
                c.font = font_sub
                c.border = thin_border

        # Total Marks Formula (summing each subject's total cell)
        tot_let = get_column_letter(col_tot_marks)
        tot_formula = "=SUM(" + ",".join(f"{ltr}{r}" for ltr in sub_total_col_letters) + ")"
        c_tot = ws.cell(r, col_tot_marks, value=tot_formula)
        c_tot.alignment = align_center
        c_tot.font = font_data
        c_tot.border = thin_border

        # Max Marks
        max_t = s.get("max_total") or (len(courses) * 100 - sum(50 for c in courses if _is_practical(c))) or 700
        max_let = get_column_letter(col_max_marks)
        c_max = ws.cell(r, col_max_marks, value=int(max_t))
        c_max.alignment = align_center
        c_max.font = font_data
        c_max.border = thin_border

        # Overall % Formula
        pct_formula = f"=ROUND({tot_let}{r}/{max_let}{r}*100,2)"
        c_pct = ws.cell(r, col_pct, value=pct_formula)
        c_pct.alignment = align_center
        c_pct.font = font_data
        c_pct.border = thin_border

        # SGPA
        c_sgpa = ws.cell(r, col_sgpa, value=s.get("sgpa") or "")
        c_sgpa.alignment = align_center
        c_sgpa.font = font_data
        c_sgpa.border = thin_border

        # CGPA
        c_cgpa = ws.cell(r, col_cgpa, value=s.get("cgpa") or "")
        c_cgpa.alignment = align_center
        c_cgpa.font = font_data
        c_cgpa.border = thin_border

        # Result
        res_val = s.get("result", "PASS").upper()
        c_res = ws.cell(r, col_result, value=res_val)
        c_res.alignment = align_center
        c_res.border = thin_border
        if res_val == "PASS":
            c_res.fill = fill_pass
            c_res.font = font_pass
        else:
            c_res.fill = fill_fail
            c_res.font = font_fail

    end_student_row = start_row + num_students - 1 if num_students else start_row

    # 6. Subject-Wise Result Analysis Table
    r_sub = end_student_row + 3
    ws.row_dimensions[r_sub].height = 16
    ws.merge_cells(f"A{r_sub}:G{r_sub}")
    t_sub = ws[f"A{r_sub}"]
    t_sub.value = "SUBJECT-WISE RESULT ANALYSIS"
    t_sub.font = font_sec_head
    t_sub.alignment = align_left

    r_sub_hdr = r_sub + 1
    ws.row_dimensions[r_sub_hdr].height = 26
    sub_headers = ["Sl No", "Course Code", "Subject Name", "Appeared", "Passed", "Failed", "Pass %"]
    for ci, h in enumerate(sub_headers, 1):
        cell = ws.cell(r_sub_hdr, ci, value=h)
        cell.fill = fill_navy
        cell.font = font_th
        cell.alignment = align_center
        cell.border = thin_border

    sub_start_data_row = r_sub_hdr + 1
    for i, course in enumerate(courses):
        curr_r = sub_start_data_row + i
        code = course["code"]
        short_code = _get_short_code(code)
        c_name = course.get("name", short_code)

        app_count = sum(1 for st in students if code in st.get("subjects", {}))
        pass_count = sum(1 for st in students if st.get("subjects", {}).get(code, {}).get("status") == "PASS")
        fail_count = sum(1 for st in students if st.get("subjects", {}).get(code, {}).get("status") == "FAIL")

        ws.cell(curr_r, 1, value=i + 1).alignment = align_center
        ws.cell(curr_r, 2, value=short_code).alignment = align_center
        ws.cell(curr_r, 3, value=c_name).alignment = align_left
        ws.cell(curr_r, 4, value=app_count).alignment = align_center
        ws.cell(curr_r, 5, value=pass_count).alignment = align_center
        ws.cell(curr_r, 6, value=fail_count).alignment = align_center

        pct_form = f"=ROUND(E{curr_r}/D{curr_r}*100,2)" if app_count else "-"
        ws.cell(curr_r, 7, value=pct_form).alignment = align_center

        for ci in range(1, 8):
            ws.cell(curr_r, ci).font = font_data
            ws.cell(curr_r, ci).border = thin_border

    # Overall Subject Summary Row
    sub_end_data_row = sub_start_data_row + num_courses - 1
    ov_row = sub_end_data_row + 1
    ws.cell(ov_row, 3, value="OVERALL (all subjects)").alignment = align_left
    ws.cell(ov_row, 3).font = font_data_bold

    tot_app = f"=SUM(D{sub_start_data_row}:D{sub_end_data_row})"
    tot_pass = f"=SUM(E{sub_start_data_row}:E{sub_end_data_row})"
    tot_fail = f"=SUM(F{sub_start_data_row}:F{sub_end_data_row})"
    ov_pct = f"=ROUND(E{ov_row}/D{ov_row}*100,2)"

    ws.cell(ov_row, 4, value=tot_app).alignment = align_center
    ws.cell(ov_row, 5, value=tot_pass).alignment = align_center
    ws.cell(ov_row, 6, value=tot_fail).alignment = align_center
    ws.cell(ov_row, 7, value=ov_pct).alignment = align_center

    for ci in range(1, 8):
        ws.cell(ov_row, ci).font = font_data_bold
        ws.cell(ov_row, ci).border = thin_border

    # 7. Overall Class Result Summary Table
    r_cls = ov_row + 3
    ws.row_dimensions[r_cls].height = 16
    ws.merge_cells(f"A{r_cls}:G{r_cls}")
    t_cls = ws[f"A{r_cls}"]
    t_cls.value = "OVERALL CLASS RESULT SUMMARY"
    t_cls.font = font_sec_head
    t_cls.alignment = align_left

    r_cls_hdr = r_cls + 1
    ws.row_dimensions[r_cls_hdr].height = 26
    cls_headers = ["Sl No", "Category", "No. of Students", "% of Total"]
    for ci, h in enumerate(cls_headers, 1):
        cell = ws.cell(r_cls_hdr, ci, value=h)
        cell.fill = fill_navy
        cell.font = font_th
        cell.alignment = align_center
        cell.border = thin_border

    # Tiers
    dist_count = sum(1 for st in students if st.get("percentage") and st["percentage"] >= 70)
    fc_count = sum(1 for st in students if st.get("percentage") and 60 <= st["percentage"] < 70)
    sc_count = sum(1 for st in students if st.get("percentage") and 35 <= st["percentage"] < 60)
    fail_35_count = sum(1 for st in students if st.get("percentage") and st["percentage"] < 35)

    tot_appeared = num_students
    tot_passed = sum(1 for st in students if st.get("result") == "PASS")
    tot_failed = sum(1 for st in students if st.get("result") != "PASS")

    tier_rows = [
        (1, "Distinction (>=70%)", dist_count),
        (2, "First Class (60-69%)", fc_count),
        (3, "Second Class (35-59%)", sc_count),
        (4, "Fail/Below 35%", fail_35_count),
    ]

    r_tier_start = r_cls_hdr + 1
    for i, (sl, cat, cnt) in enumerate(tier_rows):
        cur_r = r_tier_start + i
        ws.cell(cur_r, 1, value=sl).alignment = align_center
        ws.cell(cur_r, 2, value=cat).alignment = align_left
        ws.cell(cur_r, 3, value=cnt).alignment = align_center
        pct_expr = f"=ROUND(C{cur_r}/{tot_appeared}*100,2)" if tot_appeared else 0
        ws.cell(cur_r, 4, value=pct_expr).alignment = align_center

        for ci in range(1, 5):
            ws.cell(cur_r, ci).font = font_data
            ws.cell(cur_r, ci).border = thin_border

    # Final summary rows
    r_app = r_tier_start + len(tier_rows)
    r_pass = r_app + 1
    r_fail = r_pass + 1

    ws.cell(r_app, 2, value="Total Students Appeared").alignment = align_left
    ws.cell(r_app, 3, value=tot_appeared).alignment = align_center

    ws.cell(r_pass, 2, value="Total Passed (Overall Result)").alignment = align_left
    ws.cell(r_pass, 3, value=tot_passed).alignment = align_center
    ws.cell(r_pass, 4, value=f"=ROUND(C{r_pass}/{tot_appeared}*100,2)" if tot_appeared else 0).alignment = align_center

    ws.cell(r_fail, 2, value="Total Failed (Overall Result)").alignment = align_left
    ws.cell(r_fail, 3, value=tot_failed).alignment = align_center

    for r_idx in (r_app, r_pass, r_fail):
        for ci in range(1, 5):
            cell = ws.cell(r_idx, ci)
            cell.font = font_data_bold
            cell.border = thin_border

    # 8. Set Column Widths to match institutional layout
    ws.column_dimensions["A"].width = 6.0
    ws.column_dimensions["B"].width = 24.0
    ws.column_dimensions["C"].width = 26.0

    for i in range(num_courses):
        col1 = start_sub_col + (i * 3)
        col2 = col1 + 1
        col3 = col1 + 2
        ws.column_dimensions[get_column_letter(col1)].width = 8.0
        ws.column_dimensions[get_column_letter(col2)].width = 7.5
        ws.column_dimensions[get_column_letter(col3)].width = 10.0

    ws.column_dimensions[get_column_letter(col_tot_marks)].width = 11.0
    ws.column_dimensions[get_column_letter(col_max_marks)].width = 9.0
    ws.column_dimensions[get_column_letter(col_pct)].width = 11.0
    ws.column_dimensions[get_column_letter(col_sgpa)].width = 7.0
    ws.column_dimensions[get_column_letter(col_cgpa)].width = 7.0
    ws.column_dimensions[get_column_letter(col_result)].width = 9.0

    ws.views.sheetView[0].showGridLines = True
